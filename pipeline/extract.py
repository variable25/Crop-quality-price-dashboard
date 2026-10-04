"""
extract.py - find files in the inbox, fingerprint them, and read their rows.
"""
import hashlib
from dataclasses import dataclass

import pandas as pd
from openpyxl import load_workbook

from pipeline.clean import is_missing, normalize_key
from pipeline.config import INBOX_DIR, load_mappings

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 15        # look for the header in the first 15 rows
MIN_HEADER_MATCHES = 3       # a row needs 3+ known column names to count as the header


@dataclass
class Sheet:
    format_name: str         # "lab", "agmarknet" or "trader"
    file_type: str           # "lab" or "price"
    header_row: int          # Excel row number of the header
    rows: pd.DataFrame       # standard column names + excel_row + raw


def list_inbox_files(inbox=INBOX_DIR):
    """Excel files in the inbox, oldest first. Skips Excel's '~$' lock files."""
    if not inbox.exists():
        return []
    files = [p for p in inbox.iterdir()
             if p.is_file() and p.suffix.lower() in EXCEL_SUFFIXES and not p.name.startswith("~$")]
    return sorted(files, key=lambda p: (p.stat().st_mtime, p.name))


def file_hash(path, chunk_size=65536):
    """SHA-256 fingerprint of the file's bytes: same content -> same hash, whatever the name."""
    sha = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            sha.update(chunk)
    return sha.hexdigest()


def read_excel_rows(path):
    """Every row of the first sheet as (excel_row_number, [cell values])."""
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb.worksheets[0]
        return [(n, list(values)) for n, values in enumerate(ws.iter_rows(values_only=True), start=1)]
    finally:
        wb.close()


def find_header(rows, formats):
    """Return (position in rows, format name) of the first row that looks like a header."""
    for i, (_, values) in enumerate(rows[:HEADER_SCAN_ROWS]):
        cells = [normalize_key(v) for v in values]
        scores = {name: sum(cell in fmt["headers"] for cell in cells) for name, fmt in formats.items()}
        best = max(scores, key=scores.get)
        if scores[best] >= MIN_HEADER_MATCHES:
            return i, best
    raise ValueError(f"no header row found in the first {HEADER_SCAN_ROWS} rows")


def read_sheet(path, mappings):
    """Read one Excel file into a Sheet with standard column names."""
    rows = read_excel_rows(path)
    header_pos, format_name = find_header(rows, mappings["formats"])
    fmt = mappings["formats"][format_name]
    header_excel_row, header_values = rows[header_pos]

    original_names = ["" if is_missing(h) else str(h).strip() for h in header_values]
    standard_names = [fmt["headers"].get(normalize_key(h)) for h in header_values]

    records = []
    for excel_row, values in rows[header_pos + 1:]:
        record = {"excel_row": excel_row}
        raw = {}
        for i, value in enumerate(values):
            if i < len(original_names) and original_names[i]:
                raw[original_names[i]] = value
            if i < len(standard_names) and standard_names[i]:
                record[standard_names[i]] = value
        record["raw"] = raw
        records.append(record)

    # dtype=object: keep every value exactly as Excel gave it; cleaning happens in transform
    return Sheet(format_name, fmt["file_type"], header_excel_row, pd.DataFrame(records, dtype=object))


if __name__ == "__main__":
    mappings = load_mappings()
    files = list_inbox_files()
    print(f"{len(files)} files in {INBOX_DIR}\n")
    for path in files:
        print(path.name)
        print(f"  hash:   {file_hash(path)[:16]}...")
        try:
            sheet = read_sheet(path, mappings)
        except Exception as exc:
            print(f"  ERROR:  {type(exc).__name__}: {exc}")
        else:
            print(f"  format: {sheet.format_name} ({sheet.file_type}), "
                  f"header on Excel row {sheet.header_row}, {len(sheet.rows)} rows")
        print()